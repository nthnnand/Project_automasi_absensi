from openpyxl import load_workbook

files = [
    (r'c:\HENRI\Campus\Magang\PNS - PERHITUNGAN TERBARU JAM PULANG KHUSUS JULI27.xlsx', 'PNS'),
    (r'c:\HENRI\Campus\Magang\PPPK - PERHITUNGAN TERBARU JAM PULANG KHUSUS JULI271.xlsx', 'PPPK'),
]

for fpath, label in files:
    print('='*60)
    print(f'{label} - Sheet: Input Data Absen')
    wb = load_workbook(fpath, read_only=True, data_only=True)
    ws = wb['Input Data Absen']
    
    max_col = ws.max_column
    print(f'Max row: {ws.max_row}, Max col: {max_col}')

    # Row 3: tanggal / nomor hari
    print('\n-- Row 3 (angka/tanggal) --')
    for c in range(1, max_col+1):
        v = ws.cell(3, c).value
        if v is not None:
            print(f'  Col {c}: {repr(v)}')

    print('\n-- Row 4 (nama hari) --')
    for c in range(1, max_col+1):
        v = ws.cell(4, c).value
        if v is not None:
            print(f'  Col {c}: {repr(v)}')

    print('\n-- Row 5 (K/M/P) --')
    for c in range(1, max_col+1):
        v = ws.cell(5, c).value
        if v is not None:
            print(f'  Col {c}: {repr(v)}')

    print('\n-- Row 6 sample (first data row) --')
    for c in range(1, min(max_col+1, 30)):
        v = ws.cell(6, c).value
        if v is not None:
            print(f'  Col {c}: {repr(v)}')

    # Cek sheet Database
    if 'Database' in wb.sheetnames or 'database' in wb.sheetnames:
        sname = 'Database' if 'Database' in wb.sheetnames else 'database'
        ws2 = wb[sname]
        print(f'\n-- Sheet [{sname}] first 10 rows --')
        for r in range(1, 11):
            row = [ws2.cell(r,c).value for c in range(1, ws2.max_column+1)]
            non_none = [(i+1, v) for i,v in enumerate(row) if v is not None]
            print(f'  R{r}:', non_none[:15])

    # List all sheets
    print('\nAll sheets:', wb.sheetnames)
    wb.close()
    print()
